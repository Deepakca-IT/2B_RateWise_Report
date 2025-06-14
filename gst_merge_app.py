
import streamlit as st
import pandas as pd
from io import BytesIO
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from datetime import datetime

st.title("GSTR-2B vs GSTR-2A Merger Tool (Client Access Controlled)")

# Load access control list from GitHub
CLIENT_CSV_URL = "https://raw.githubusercontent.com/Deepakca-IT/2B_RateWise_Report/main/clients.csv"

def get_authorized_clients():
    return pd.read_csv(CLIENT_CSV_URL)

def is_authorized(gstin, access_df):
    today = pd.Timestamp.today().normalize()
    match = access_df[access_df['GSTIN'] == gstin]
    if not match.empty:
        start = pd.to_datetime(match.iloc[0]['Start Date'])
        end = pd.to_datetime(match.iloc[0]['End Date'])
        return start <= today <= end
    return False

def extract_gstin_from_readme(file):
    try:
        readme = pd.read_excel(file, sheet_name="Read me", header=None)
        gstin = str(readme.iloc[5, 2]).strip()
        return gstin
    except Exception:
        st.error("❌ Could not read GSTIN from the 'Read me' sheet.")
        return None

def load_gstr2b(file):
    columns = [
        "GSTIN of supplier", "Trade/Legal name", "Invoice number", "Invoice type",
        "Invoice Date", "Invoice Value(₹)", "Place of supply", "Supply Attract Reverse Charge",
        "Taxable Value (₹)", "Integrated Tax(₹)", "Central Tax(₹)",
        "State/UT Tax(₹)", "Cess(₹)", "GSTR-1/5 Period", "GSTR-1/5 Filing Date",
        "ITC Availability", "Reason", "Applicable % of Tax Rate", "Source", "IRN",
        "IRN Date"
    ]
    df = pd.read_excel(file, sheet_name="B2B", header=None, skiprows=6)
    df.columns = columns
    df['Invoice Date'] = pd.to_datetime(df['Invoice Date'], dayfirst=True, errors='coerce')
    df['Month-Year'] = df['Invoice Date'].dt.strftime('%B-%y')
    return df

def load_and_clean_gstr2a(files):
    all_data = []
    columns = [
        "GSTIN of supplier", "Trade/Legal name of the Supplier", "Invoice number", "Invoice type",
        "Invoice Date", "Invoice Value (₹)", "Place of supply", "Supply Attract Reverse Charge",
        "Rate (%)", "Taxable Value (₹)", "Integrated Tax  (₹)", "Central Tax (₹)",
        "State/UT tax (₹)", "Cess  (₹)", "GSTR-1/5 Filing Status", "GSTR-1/5 Filing Date",
        "GSTR-1/5 Filing Period", "GSTR-3B Filing Status", "Amendment made, if any",
        "Tax Period in which Amended", "Effective date of cancellation", "Source","IRN","IRN date"
    ]
    for file in files:
        df = pd.read_excel(file, sheet_name="B2B", header=None, skiprows=6)
        df.columns = columns
        df = df[~df["Invoice number"].astype(str).str.contains("-Total", na=False)]
        df = df[df["Invoice number"].notna() & df["GSTIN of supplier"].notna()]
        df['Invoice Date'] = pd.to_datetime(df['Invoice Date'], dayfirst=True, errors='coerce')
        df.rename(columns={
            "Trade/Legal name of the Supplier": "Trade/Legal name",
            "Rate (%)": "Rate(%)",
            "Taxable Value (₹)": "Taxable Value (₹)",
            "Integrated Tax  (₹)": "Integrated Tax(₹)",
            "Central Tax (₹)": "Central Tax(₹)",
            "State/UT tax (₹)": "State/UT Tax(₹)",
            "Cess  (₹)": "Cess(₹)"
        }, inplace=True)
        all_data.append(df)
    return pd.concat(all_data, ignore_index=True)

def prepare_output_excel(df):
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Merged Data")
        ws = writer.sheets["Merged Data"]
        last_col = get_column_letter(df.shape[1])
        table = Table(displayName="MergedTable", ref=f"A1:{last_col}{df.shape[0]+1}")
        style = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True, showColumnStripes=True)
        table.tableStyleInfo = style
        ws.add_table(table)
        for col in ws.columns:
            max_len = max(len(str(cell.value)) if cell.value else 0 for cell in col)
            ws.column_dimensions[col[0].column_letter].width = max_len + 2
    output.seek(0)
    return output

gstr2b_file = st.file_uploader("Upload GSTR-2B Excel File (B2B Sheet)", type="xlsx")

if gstr2b_file:
    gstin_in_file = extract_gstin_from_readme(gstr2b_file)
    access_df = get_authorized_clients()

    if gstin_in_file and is_authorized(gstin_in_file, access_df):
        st.success(f"✅ Authorized GSTIN: {gstin_in_file}")
        df2b = load_gstr2b(gstr2b_file)
        available_months = sorted(df2b['Month-Year'].dropna().unique().tolist())

        st.markdown("### Available Months in GSTR-2B:")
        st.write(available_months)

        gstr2a_files = st.file_uploader("Upload GSTR-2A Excel File(s)", type="xlsx", accept_multiple_files=True)

        if gstr2a_files:
            df2a = load_and_clean_gstr2a(gstr2a_files)

            if st.button("Reconcile Now"):
                with st.spinner("Reconciling..."):
                    merged = pd.merge(
                        df2b,
                        df2a[["Invoice number", "GSTIN of supplier", "Rate(%)", "Taxable Value (₹)",
                              "Integrated Tax(₹)", "Central Tax(₹)", "State/UT Tax(₹)", "Cess(₹)"]],
                        on=["Invoice number", "GSTIN of supplier"],
                        how="left"
                    )

                    to_drop = [col for col in merged.columns if col.endswith("_x")]
                    merged.drop(columns=to_drop, inplace=True)

                    insert_cols = [col for col in merged.columns if col.endswith("_y")]
                    pos = merged.columns.get_loc("Supply Attract Reverse Charge") + 1
                    cols = merged.columns.tolist()
                    for col in reversed(insert_cols):
                        cols.remove(col)
                        cols.insert(pos, col)
                    merged = merged[cols]

                    excel_bytes = prepare_output_excel(merged)

                    # Calculate and display taxable value summary
                    taxable_2b = df2b["Taxable Value (₹)"].sum()
                    taxable_final = merged["Taxable Value (₹)_y"].sum()
                    
                    st.markdown("### 📊 Taxable Value Comparison")
                    st.write(f"**Taxable Value as per GSTR-2B:** ₹{taxable_2b:,.2f}")
                    st.write(f"**Taxable Value (from GSTR-2A) in Final Report:** ₹{taxable_final:,.2f}")
                    st.info("Compare the values above to ensure the reconciliation is aligned.")

                    st.success("Merge completed successfully!")
                    st.download_button("Download Merged Excel", data=excel_bytes,
                                       file_name="merged_gstr_data.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    else:
        st.error("❌ This GSTIN is not authorized or access period has expired.")
